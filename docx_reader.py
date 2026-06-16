import os
import csv
import pandas as pd
from docx import Document


def read_docx(file_path: str) -> str:
    doc = Document(file_path)
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def read_csv_wi(file_path: str) -> str:
    lines = []
    for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            with open(file_path, newline="", encoding=enc) as f:
                reader = csv.DictReader(f)
                for i, row in enumerate(reader, 1):
                    parts = [f"STEP {i}"]
                    for k, v in row.items():
                        if v and str(v).strip():
                            parts.append(f"  {k}: {v.strip()}")
                    lines.append("\n".join(parts))
            break
        except UnicodeDecodeError:
            continue
    return "\n\n".join(lines)


def read_xlsx_wi(file_path: str) -> str:
    """
    Read a formatted Excel Work Instruction.
    These are typically layout/print documents — NOT data tables.
    We extract all non-empty cell values row by row and build
    a clean readable text block for the LLM.
    """
    # Read without headers since the sheet is a formatted doc, not a table
    df = pd.read_excel(file_path, header=None)
    df = df.fillna("")

    sections = []
    current_section = []

    for _, row in df.iterrows():
        # Collect all non-empty, non-duplicate cell values from this row
        seen = set()
        vals = []
        for v in row.values:
            text = str(v).strip()
            if text and text not in seen:
                seen.add(text)
                vals.append(text)

        if not vals:
            # Blank row = section break
            if current_section:
                sections.append("\n".join(current_section))
                current_section = []
        else:
            current_section.append(" | ".join(vals))

    if current_section:
        sections.append("\n".join(current_section))

    return "\n\n".join(sections)


def read_wi_file(file_path: str) -> str:
    """Auto-detect file type and read accordingly."""
    ext = os.path.splitext(file_path)[-1].lower()
    if ext == ".docx":
        return read_docx(file_path)
    elif ext == ".csv":
        return read_csv_wi(file_path)
    elif ext in (".xlsx", ".xls"):
        return read_xlsx_wi(file_path)
    else:
        raise ValueError(f"Unsupported WI file type '{ext}'. Use .docx, .xlsx, or .csv.")